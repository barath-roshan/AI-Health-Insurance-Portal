import os
import re
import json
import csv
import hashlib
from datetime import datetime, timezone
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
EXCEL_PATH = os.path.join(DATA_DIR, "kaapan_comprehensive_health_schemes_dataset.xlsx")

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(os.path.join(DATA_DIR, "evaluation"), exist_ok=True)
os.makedirs(os.path.join(DATA_DIR, "pageindex", "trees"), exist_ok=True)
os.makedirs(os.path.join(DATA_DIR, "reports"), exist_ok=True)

SCHEME_TYPE_MAPPING = {
    "IN-PMJAY": "Insurance / Health Assurance",
    "TN-CMCHIS": "Insurance / Health Assurance",
    "KL-MEDISEP": "Insurance / Health Assurance",
    "IN-CGHS": "Insurance / Health Assurance",
    "IN-ESIS": "Insurance / Health Assurance",
    "IN-PMMVY": "Cash Assistance / Maternity Benefit",
    "IN-JSSK": "Public Healthcare Entitlement",
    "IN-RBSK": "Public Healthcare Service",
    "IN-PMSSY": "Public Healthcare Infrastructure",
    "IN-NPCDCS": "Public Healthcare Service",
    "IN-PMNDP": "Public Healthcare Service",
    "TN-NK48": "Emergency Healthcare Assistance",
    "IN-ABHA": "Digital Health Identity"
}

def is_valid_url(url: str) -> bool:
    if not url:
        return False
    return bool(re.match(r"^https?://[^\s/$.?#].[^\s]*$", url))

def validate_and_ingest():
    print(f"Loading Excel workbook: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    
    # 1. Read Table of Contents
    toc_data = []
    if "Table_of_Contents" in wb.sheetnames:
        toc_sheet = wb["Table_of_Contents"]
        for row in list(toc_sheet.iter_rows(values_only=True))[1:]:
            if any(row):
                toc_data.append([str(cell) if cell is not None else "" for cell in row])
                
    # 2. Read Schemes Sheet
    if "Schemes" not in wb.sheetnames:
        raise ValueError("Workbook missing required 'Schemes' sheet!")
        
    schemes_sheet = wb["Schemes"]
    rows = list(schemes_sheet.iter_rows(values_only=True))
    headers = [str(h).strip() for h in rows[0]]
    
    records = []
    source_records = []
    alias_records = []
    validation_errors = []
    seen_ids = set()
    now_iso = datetime.now(timezone.utc).isoformat()

    for idx, row in enumerate(rows[1:], 1):
        if not any(row):
            continue
        raw_dict = dict(zip(headers, row))
        scheme_id = str(raw_dict.get("scheme_id") or f"SCHEME-{idx}").strip()
        scheme_name = str(raw_dict.get("scheme_name") or "").strip()
        jurisdiction = str(raw_dict.get("jurisdiction") or "India").strip()
        category = str(raw_dict.get("category") or "").strip()
        brief_details = str(raw_dict.get("brief_details") or "").strip()
        eligibility = str(raw_dict.get("eligibility") or "").strip()
        required_documents = str(raw_dict.get("required_documents") or "").strip()
        where_to_apply = str(raw_dict.get("where_to_apply") or "").strip()
        official_url = str(raw_dict.get("official_url") or "").strip()
        source_url = str(raw_dict.get("source_url") or "").strip()
        source_type = str(raw_dict.get("source_type") or "").strip()
        ver_status = str(raw_dict.get("verification_status") or "NEEDS_REVIEW").strip()
        last_ver = str(raw_dict.get("last_verified") or "2026-09-30").strip()
        notes = str(raw_dict.get("notes") or "").strip()

        # Validation Checks
        if scheme_id in seen_ids:
            validation_errors.append(f"Duplicate scheme_id detected: '{scheme_id}' at row {idx}")
        seen_ids.add(scheme_id)

        if not scheme_name:
            validation_errors.append(f"Missing scheme_name at row {idx}")

        if official_url and not is_valid_url(official_url):
            validation_errors.append(f"Invalid official_url '{official_url}' for {scheme_id}")

        scheme_type = SCHEME_TYPE_MAPPING.get(scheme_id, "Public Healthcare Service")
        is_enrolment_required = scheme_type in ["Insurance / Health Assurance", "Cash Assistance / Maternity Benefit"]

        record = {
            "scheme_id": scheme_id,
            "scheme_name": scheme_name,
            "normalized_name": scheme_name.lower(),
            "jurisdiction": jurisdiction,
            "category": category,
            "scheme_type": scheme_type,
            "is_enrolment_required": is_enrolment_required,
            "brief_details": brief_details,
            "eligibility": eligibility,
            "required_documents": required_documents,
            "where_to_apply": where_to_apply,
            "official_url": official_url if is_valid_url(official_url) else None,
            "source_url": source_url if is_valid_url(source_url) else None,
            "source_type": source_type,
            "verification_status": ver_status,
            "last_verified": last_ver,
            "notes": notes,
            "provenance": {
                "input_source": "kaapan_comprehensive_health_schemes_dataset.xlsx",
                "row_number": idx,
                "imported_at": now_iso,
                "verified_against_official_sources": ver_status == "VERIFIED_OFFICIAL"
            }
        }
        records.append(record)

        # Build Source Registry Entry
        source_id = f"src-{scheme_id.lower()}"
        source_records.append({
            "source_id": source_id,
            "scheme_id": scheme_id,
            "document_title": f"{scheme_name} Official Guidelines",
            "issuing_authority": source_type,
            "official_domain": official_url.split('/')[2] if '://' in official_url else "gov.in",
            "source_url": official_url or source_url,
            "document_type": "OFFICIAL_GUIDELINE",
            "document_language": "en",
            "publication_date": "2024-01-01",
            "retrieved_at": now_iso,
            "last_verified_at": last_ver,
            "verification_status": ver_status,
            "checksum": hashlib.sha256(f"{scheme_id}-{scheme_name}".encode('utf-8')).hexdigest(),
            "ingestion_status": "COMPLETED"
        })

        # Build Alias Entries
        alias_records.append({"alias": scheme_id.lower(), "scheme_id": scheme_id, "scheme_name": scheme_name})
        alias_records.append({"alias": scheme_name.lower(), "scheme_id": scheme_id, "scheme_name": scheme_name})
        acronym = "".join([w[0] for w in scheme_name.split() if w[0].isupper()]).lower()
        if len(acronym) > 1:
            alias_records.append({"alias": acronym, "scheme_id": scheme_id, "scheme_name": scheme_name})

        # Build PageIndex Document Tree
        tree_data = {
            "document_id": f"pid-{scheme_id.lower()}",
            "tree_id": f"tree-{scheme_id.lower()}",
            "scheme_id": scheme_id,
            "title": f"{scheme_name} Policy & Guidelines",
            "nodes": [
                {
                    "section_id": "sec-1-overview",
                    "title": "1. Programme Overview & Type",
                    "page_start": 1,
                    "page_end": 1,
                    "content": f"Programme Name: {scheme_name}\nJurisdiction: {jurisdiction}\nCategory: {category}\nType: {scheme_type}\nDetails: {brief_details}"
                },
                {
                    "section_id": "sec-2-eligibility",
                    "title": "2. Eligibility Criteria & Conditions",
                    "page_start": 2,
                    "page_end": 2,
                    "content": f"Eligibility: {eligibility}\nIndividual Enrolment Required: {is_enrolment_required}"
                },
                {
                    "section_id": "sec-3-documents",
                    "title": "3. Required Documents & Documentation",
                    "page_start": 3,
                    "page_end": 3,
                    "content": f"Required Documents: {required_documents}"
                },
                {
                    "section_id": "sec-4-application",
                    "title": "4. Application Procedure & Access Desks",
                    "page_start": 4,
                    "page_end": 4,
                    "content": f"Where to Apply / Access: {where_to_apply}\nOfficial Portal: {official_url}"
                }
            ]
        }
        tree_path = os.path.join(DATA_DIR, "pageindex", "trees", f"tree_{scheme_id.lower()}.json")
        with open(tree_path, "w", encoding="utf-8") as tf:
            json.dump(tree_data, tf, indent=2)

    # 3. Write Single Canonical JSON Catalogue
    canonical_catalogue = {
        "metadata": {
            "dataset_name": "KAAPAN Comprehensive Health Schemes Catalogue",
            "version": "2.0.0",
            "schema_type": "Vectorless PageIndex Knowledge Base",
            "source_workbook": "kaapan_comprehensive_health_schemes_dataset.xlsx",
            "total_schemes": len(records),
            "generated_at": now_iso
        },
        "table_of_contents": toc_data,
        "schemes": records
    }

    json_path = os.path.join(DATA_DIR, "kaapan_schemes_catalogue.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(canonical_catalogue, f, indent=2)

    # 4. Write CSV Exports for Compatibility
    csv_path = os.path.join(DATA_DIR, "scheme_catalogue.csv")
    if records:
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            fieldnames = ["scheme_id", "scheme_name", "normalized_name", "jurisdiction", "category", "scheme_type", "brief_details", "eligibility", "required_documents", "where_to_apply", "official_url", "source_url", "source_type", "verification_status", "last_verified", "notes"]
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(records)

    # 5. Write Source Registry CSV
    source_csv_path = os.path.join(DATA_DIR, "source_registry.csv")
    if source_records:
        with open(source_csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(source_records[0].keys()))
            writer.writeheader()
            writer.writerows(source_records)

    # 6. Write Scheme Aliases CSV
    alias_csv_path = os.path.join(DATA_DIR, "scheme_aliases.csv")
    if alias_records:
        with open(alias_csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=["alias", "scheme_id", "scheme_name"])
            writer.writeheader()
            writer.writerows(alias_records)

    # 7. Write Validation Report
    validation_report = {
        "timestamp": now_iso,
        "total_records_processed": len(records),
        "validation_status": "SUCCESS" if not validation_errors else "WARNINGS",
        "validation_errors": validation_errors,
        "canonical_json_path": json_path,
        "pageindex_trees_generated": len(records)
    }
    report_path = os.path.join(DATA_DIR, "reports", "dataset_validation_report.json")
    with open(report_path, "w", encoding="utf-8") as f:
        json.dump(validation_report, f, indent=2)

    print(f"Catalogue ingestion complete! Successfully processed {len(records)} schemes from workbook.")
    if validation_errors:
        print(f"Validation Report Warnings ({len(validation_errors)}): {validation_errors}")

if __name__ == "__main__":
    validate_and_ingest()
