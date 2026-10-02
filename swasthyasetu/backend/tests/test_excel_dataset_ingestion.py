import os
import json
import pytest
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
EXCEL_PATH = os.path.join(DATA_DIR, "kaapan_comprehensive_health_schemes_dataset.xlsx")
CANONICAL_JSON_PATH = os.path.join(DATA_DIR, "kaapan_schemes_catalogue.json")

def test_excel_file_exists():
    assert os.path.exists(EXCEL_PATH), f"Excel dataset workbook not found at {EXCEL_PATH}"

def test_workbook_sheets():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    assert "Table_of_Contents" in wb.sheetnames
    assert "Schemes" in wb.sheetnames
    
    sheet = wb["Schemes"]
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) >= 14  # Header + 13 schemes

def test_canonical_json_catalogue():
    assert os.path.exists(CANONICAL_JSON_PATH), f"Canonical JSON dataset missing at {CANONICAL_JSON_PATH}"
    with open(CANONICAL_JSON_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)

    assert "metadata" in data
    assert "schemes" in data
    schemes = data["schemes"]
    assert len(schemes) == 13

    scheme_ids = [s["scheme_id"] for s in schemes]
    expected_ids = [
        "IN-PMJAY", "TN-CMCHIS", "IN-ESIC", "IN-CGHS", "IN-RAN",
        "IN-JSY", "IN-JSSK", "IN-RBSK", "IN-PMSSY", "IN-NPCDCS",
        "IN-PMNDP", "TN-NK48", "IN-ABHA"
    ]
    for eid in expected_ids:
        assert eid in scheme_ids, f"Expected scheme_id '{eid}' missing from catalogue"

def test_scheme_categorization_and_type():
    with open(CANONICAL_JSON_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)
    
    schemes_dict = {s["scheme_id"]: s for s in data["schemes"]}
    
    # ABHA must be Digital Health Identity (not insurance)
    abha = schemes_dict.get("IN-ABHA")
    assert abha["scheme_type"] == "Digital Health Identity"
    assert abha["is_enrolment_required"] is False

    # PMSSY must be Healthcare Infrastructure
    pmssy = schemes_dict.get("IN-PMSSY")
    assert pmssy["scheme_type"] == "Public Healthcare Infrastructure"
    assert pmssy["is_enrolment_required"] is False

    # PMJAY must be Insurance / Health Assurance
    pmjay = schemes_dict.get("IN-PMJAY")
    assert pmjay["scheme_type"] == "Insurance / Health Assurance"

def test_pageindex_trees_generated():
    tree_dir = os.path.join(DATA_DIR, "pageindex", "trees")
    expected_ids = [
        "in-pmjay", "tn-cmchis", "in-esic", "in-cghs", "in-ran",
        "in-jsy", "in-jssk", "in-rbsk", "in-pmssy", "in-npcdcs",
        "in-pmndp", "tn-nk48", "in-abha"
    ]
    for eid in expected_ids:
        tree_file = os.path.join(tree_dir, f"tree_{eid}.json")
        assert os.path.exists(tree_file), f"PageIndex section tree missing for '{eid}'"
        with open(tree_file, "r", encoding="utf-8") as f:
            tdata = json.load(f)
            assert "nodes" in tdata
            assert len(tdata["nodes"]) >= 4
